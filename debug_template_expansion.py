#!/usr/bin/env python3
"""
Debug script to test template expansion and placeholder creation
"""

import os
import sys
from pathlib import Path

# Add the src directory to the path
sys.path.insert(0, str(Path(__file__).parent / "src"))

from src.core.generation.template_processor import TemplateProcessor
from docx import Document
from io import BytesIO

def test_template_expansion():
    """Test template expansion to see what placeholders are created."""
    
    print("Testing template expansion...")
    
    try:
        # Create a template processor
        tp = TemplateProcessor(template_type='vertical', font_scheme='Arial')
        
        # Check if the expanded template buffer exists
        if hasattr(tp, '_expanded_template_buffer') and tp._expanded_template_buffer:
            print("✅ Template expansion successful")
            
            # Load the expanded template
            if hasattr(tp._expanded_template_buffer, 'seek'):
                tp._expanded_template_buffer.seek(0)
                doc = Document(tp._expanded_template_buffer)
            else:
                print("❌ Expanded template buffer is not seekable")
                return
            
            # Check what placeholders are in the expanded template
            print("\nChecking expanded template for placeholders:")
            placeholders_found = set()
            
            for table_idx, table in enumerate(doc.tables):
                print(f"\nTable {table_idx + 1}:")
                for row_idx, row in enumerate(table.rows):
                    for col_idx, cell in enumerate(row.cells):
                        cell_text = cell.text.strip()
                        if cell_text:
                            print(f"  Cell ({row_idx + 1}, {col_idx + 1}): '{cell_text}'")
                            
                            # Find placeholders in this cell
                            import re
                            placeholder_matches = re.findall(r'\{\{.*?\}\}', cell_text)
                            for match in placeholder_matches:
                                placeholders_found.add(match)
                                print(f"    → Found placeholder: {match}")
            
            print(f"\nSummary of placeholders found:")
            for placeholder in sorted(placeholders_found):
                print(f"  - {placeholder}")
            
            # Check if DescAndWeight placeholder is present
            if any('DescAndWeight' in p for p in placeholders_found):
                print("✅ DescAndWeight placeholder found - preroll descriptions should work")
            else:
                print("❌ DescAndWeight placeholder NOT found - this is why preroll descriptions aren't working")
                
        else:
            print("❌ Template expansion failed - no expanded template buffer")
            
    except Exception as e:
        print(f"❌ Error during testing: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    test_template_expansion() 