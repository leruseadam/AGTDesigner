#!/usr/bin/env python3
"""
Test script to verify that brand placeholders are now showing in double template output.
"""

import sys
import os
sys.path.append(os.path.dirname(os.path.abspath(__file__)))

from src.core.generation.template_processor import TemplateProcessor
from src.core.generation.unified_font_sizing import get_font_size_by_marker
from docx import Document
from docx.shared import Pt
import tempfile
import os

def test_double_template_brand_placeholder():
    """Test that double template now shows brand placeholders."""
    print("Testing Double Template Brand Placeholder Fix")
    
    # Create test record with brand information
    test_record = {
        'ProductName': 'Test Product',
        'ProductBrand': 'Test Brand',
        'ProductStrain': 'Test Strain',
        'ProductType': 'flower',
        'ProductVendor': 'Test Vendor',
        'Lineage': 'HYBRID',
        'Price': '$25.00',
        'WeightUnits': '3.5g',
        'Description': 'Test Description'
    }
    
    try:
        # Initialize template processor for double template
        tp = TemplateProcessor('double', {})
        
        # Process the test record
        documents = tp.process_records([test_record])
        
        if not documents:
            print("✗ No documents generated")
            return False
        
        # Save the first document to a temporary file
        doc = documents[0] if isinstance(documents, list) else documents
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp_file:
            doc.save(tmp_file.name)
            tmp_path = tmp_file.name
        
        # Check if brand content is present in the document
        brand_found = False
        brand_marker_found = False
        
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        text = paragraph.text
                        if 'Test Brand' in text:
                            brand_found = True
                            print(f"✓ Found brand content: '{text.strip()}'")
                        if 'PRODUCTBRAND_CENTER_START' in text or 'PRODUCTBRAND_CENTER_END' in text:
                            brand_marker_found = True
                            print(f"✓ Found brand markers in: '{text.strip()}'")
        
        # Clean up temporary file
        os.unlink(tmp_path)
        
        if brand_found:
            print("✓ Brand content is now showing in double template output")
            return True
        else:
            print("✗ Brand content is still not showing in double template output")
            return False
            
    except Exception as e:
        print(f"✗ Error testing double template brand placeholder: {e}")
        return False

if __name__ == "__main__":
    print("Testing Double Template Brand Placeholder Fix")
    print("=" * 50)
    
    success = test_double_template_brand_placeholder()
    
    if success:
        print("\n✓ Double template brand placeholder fix test PASSED")
    else:
        print("\n✗ Double template brand placeholder fix test FAILED") 